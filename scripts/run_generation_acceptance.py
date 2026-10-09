"""Acceptance thật trên mock cục bộ; không model/provider, không giả lập browser.

uv run python scripts/run_generation_acceptance.py --out artifacts/generation-acceptance
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
import sys
import threading
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
import yaml

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from qctool.executability import analyze_executability
from qctool.pipeline import PipelineConfig, PipelineResult, run_pipeline
from tests.mock_server import create_server

SAMPLE = REPO / "samples/acceptance/generation"


def verify_acceptance(result: PipelineResult) -> dict:
    """Đối chiếu nguồn, audit, IR, runner và các artifact đã sinh thật."""
    assert result.ok, result.error or "Pipeline không hoàn tất"
    assert result.run_result is not None
    assert not result.generation_success, "Provisioning chưa giải được phải ngăn coverage success"
    requirements = yaml.safe_load(result.requirements_path.read_text(encoding="utf-8"))
    expected = {"US-acceptance-001", "AC-acceptance-001-1", "FR-acceptance-001"}
    assert {r["id"] for r in requirements["requirements"]} == expected
    assert all(r["source"]["file"] and r["source"]["line"] for r in requirements["requirements"])
    flow = requirements["flows"]["flow"][0]
    assert flow["steps"][0]["action"] == "Open /ui/du-an"
    assert flow["steps"][0]["expected"] == "URL /ui/du-an"

    packet = json.loads(result.generation_path.read_text(encoding="utf-8"))
    ir = yaml.safe_load(result.testcases_path.read_text(encoding="utf-8"))
    generation = result.generation_result
    assert packet['status'] == result.generation_status and packet['success'] is False
    accepted = generation.accepted_generated_cases
    assert len(accepted) == 1
    generated = accepted[0]
    assert analyze_executability(generated).executable
    assert flow['id'] in generated.traces_to
    record = next(h for h in generation.agent_history if h.case.id == generated.id)
    assert record.dry_run_passed
    assert generated.source_intent['expected'] == flow['steps'][0]['expected']
    obligation_id = generated.obligation_ids[0]

    results = {case.case_id: case for case in result.run_result.cases}
    emitted = {case["id"]: case for case in ir["testcases"]}
    assert set(results) == set(emitted) == {case["id"] for case in packet["cases"]}
    assert packet["cases"] == ir["testcases"]
    assert all(case.get("status") != "Pass" for case in emitted.values())
    execution = results[generated.id]
    assert execution.status == "Pass" and not execution.manual
    assert [(s.action, s.ok) for s in execution.steps] == [
        ("ui.session", True), ("ui.goto", True), ("assert.url_matches", True)]
    assert "/ui/du-an" in execution.steps[1].detail
    assert execution.evidence.get("trace")
    assert (result.out_dir / "evidence" / execution.evidence["trace"]).is_file()
    for case in results.values():
        if case.status == "Pass":
            assert case.steps and all(step.ok for step in case.steps)
            assert len(case.steps) == len(emitted[case.case_id]["steps"])
            assert emitted[case.case_id].get("execution", "auto") == "auto" and not case.manual
        if emitted[case.case_id].get("execution", "auto") == "manual":
            assert case.status == "Not Executed" and case.steps == []
    assert result.counts["Not Executed"] > 0
    assert result.counts["Fail"] == result.counts["Blocked"] == 0

    disk = json.loads((result.run_dir / "results.json").read_text(encoding="utf-8"))
    assert disk == result.run_result.to_dict()
    wb = load_workbook(result.xlsx_path)
    ws = wb["Test Cases"]
    rows = {ws.cell(row, 1).value: row for row in range(2, ws.max_row + 1)
            if ws.cell(row, 1).value}
    assert set(rows) == set(results)
    for id, case in results.items():
        assert ws.cell(rows[id], 10).value == case.status
        assert ws.cell(rows[id], 9).value and ws.cell(rows[id], 12).value
    wb.close()
    with (result.run_dir / "results.csv").open(encoding="utf-8-sig", newline="") as stream:
        csv_results = {row["Mã test"]: row["Kết quả"] for row in csv.DictReader(stream)}
    assert csv_results == {id: case.status for id, case in results.items()}
    junit = ET.parse(result.run_dir / "junit.xml").getroot()
    assert int(junit.attrib["tests"]) == len(results)
    assert int(junit.attrib["skipped"]) == result.counts["Not Executed"]
    assert int(junit.attrib["failures"]) == int(junit.attrib["errors"]) == 0
    assert {case.attrib["name"] for case in junit.findall("testcase")} == set(results)
    for case in junit.findall("testcase"):
        assert (case.find("skipped") is not None) == (
            results[case.attrib["name"]].status == "Not Executed")
    markdown = result.report_path.read_text(encoding="utf-8")
    assert generated.id in markdown  # evidence section links the real browser trace
    assert all(id in markdown for id, case in results.items() if case.manual)
    for row in (f"| Tổng số test | {len(results)} |", f"| Đạt | {result.counts['Pass']} |",
                f"| Chưa chạy | {result.counts['Not Executed']} |",
                f"| Bị chặn | {result.counts['Blocked']} |",
                f"| **Lỗi** | **{result.counts['Fail']}** |"):
        assert row in markdown
    return {
        "pipeline_ok": result.ok, "generation_success": result.generation_success,
        "generation_status": result.generation_status,
        "requirements": sorted(expected), "generated_executable_case": generated.id,
        "obligation": obligation_id, "source": generated.source_ref,
        "counts": result.counts, "cases": {id: c.status for id, c in results.items()},
        "generation_passes": len(generation.history),
        "artifacts": [str(path) for path in (
            result.requirements_path, result.out_dir / "app-model.yaml", result.generation_path,
            result.testcases_path, result.xlsx_path, result.run_dir / "results.json",
            result.run_dir / "results.csv", result.run_dir / "junit.xml", result.report_path)],
        "verified": True,
    }


class RouteReplay:
    """Recorded tool fixture for a declared route. Never a live model claim."""
    def __init__(self, route):
        self.route = route
    def chat(self, messages, tools, **kwargs):
        if tools[0]['function']['name'] == 'design_tests':
            intents = json.loads(messages[-1]['content'])
            name, args = 'design_tests', {'intents':[{'id':i['id'], 'title':i['title'], 'source_quote':i['expected']} for i in intents]}
        else:
            intent = json.loads(messages[1]['content'])['intent']
            if intent['expected'] != 'URL ' + self.route:
                return {'role':'assistant','content':'No recorded safe draft for this source','tool_calls':[]}
            name, args = 'draft_case', {'title':'Declared route', 'steps':[
                {'seq':1,'kind':'given','action':'ui.session','args':{'mode':'anonymous'}},
                {'seq':2,'kind':'when','action':'ui.goto','args':{'url':self.route}},
                {'seq':3,'kind':'then','action':'assert.url_matches','args':{'pattern':self.route+'$'},
                 'source':{'req_id':intent['req_id'],'quote':intent['expected']}}]}
        return {'role':'assistant','content':None,'tool_calls':[{'id':name,'type':'function','function':{'name':name,'arguments':json.dumps(args)}}]}


def run_acceptance(out_dir: Path, on_progress=None) -> tuple[PipelineResult, dict]:
    server = create_server(0, "/api/v1")
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    host, port = server.server_address[:2]
    base = f"http://{host}:{port}"
    try:
        result = run_pipeline(PipelineConfig(
            docs_dir=SAMPLE, openapi=SAMPLE / "openapi.yaml", feature="acceptance",
            module="ACCEPTANCE", web_url=base, base_url=base + "/api/v1",
            auth_mode="manual", timeout_seconds=8, out_dir=out_dir,
            ui_evidence="always", ui_trace=True), on_progress=on_progress, resolution_llm=RouteReplay("/ui/du-an"))
        checks = verify_acceptance(result)
        (out_dir / "acceptance-checks.json").write_text(
            json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
        return result, checks
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Acceptance generation trên mock + Chromium thật")
    parser.add_argument("--out", type=Path, default=REPO / "artifacts/generation-acceptance")
    args = parser.parse_args(argv)
    try:
        _, checks = run_acceptance(args.out, lambda stage, message: print(f"[{stage}] {message}"))
    except Exception as exc:
        print(f"Acceptance failed: {exc}", file=sys.stderr)
        return 1
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
