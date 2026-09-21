"""Sinh file Excel template testcase từ spec 14 field của team.

Nguồn spec: ảnh `Screenshot From 2026-09-17 10-28-26.png` — bảng
"Standard Test Case Template Structure", 2 cột `Field Name` | `Description / Example Entry`.

Chạy:   uv run scripts/build_testcase_template.py
Output: templates/testcase/testcase-template.xlsx

Quy ước quan trọng: **một dòng = một testcase** (không phải một bước).
Các bước nằm trong cùng một ô, đánh số `1.` `2.` `3.` — đúng như ảnh spec mô tả.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "templates" / "testcase" / "testcase-template.xlsx"

# ── Bảng màu ────────────────────────────────────────────────────────────────
GEN_HDR_FILL = PatternFill("solid", fgColor="1F4E79")   # cột tool tự điền
RUN_HDR_FILL = PatternFill("solid", fgColor="9C5700")   # cột điền khi chạy
GEN_HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
BODY_FONT = Font(size=10)
ID_FONT = Font(size=10, bold=True, color="1F4E79")
HINT_FONT = Font(size=9, italic=True, color="808080")

THIN = Side(style="thin", color="D1D1D1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
GREY = PatternFill("solid", fgColor="E7E6E6")

# ── 14 cột, ĐÚNG thứ tự trong ảnh spec ─────────────────────────────────────
# (tên cột, bề rộng, nhóm, mô tả để hiện khi hover)
COLUMNS = [
    ("Test Case ID", 15, "gen",
     "Mã duy nhất cho mỗi testcase. Ví dụ: TC_AUTH_001"),
    ("Module / Feature Name", 24, "gen",
     "Module hoặc tính năng đang được test. Ví dụ: User Authentication - Login"),
    ("Test Title / Summary", 38, "gen",
     "Mô tả ngắn gọn testcase kiểm tra cái gì. Ví dụ: Verify login with valid credentials"),
    ("Description / Objective", 34, "gen",
     "Mục tiêu của testcase. Ví dụ: Ensure registered users can access the dashboard"),
    ("Preconditions", 30, "gen",
     "Điều kiện phải có TRƯỚC khi chạy test. Ví dụ: User account exists and is active"),
    ("Test Data", 26, "gen",
     "Dữ liệu đầu vào cụ thể. Ví dụ: Username: user@test.com, Password: Secure123!"),
    ("Test Steps", 48, "gen",
     "Các bước thực hiện, đánh số 1. 2. 3. trong cùng một ô"),
    ("Expected Result", 36, "gen",
     "Kết quả hệ thống PHẢI trả ra. Viết cụ thể, kiểm chứng được"),
    ("Actual Result", 30, "run",
     "Điền khi chạy: hệ thống thực tế đã làm gì"),
    ("Status", 16, "run",
     "Pass / Fail / Blocked / Not Executed — chọn từ dropdown"),
    ("Priority / Severity", 16, "gen",
     "High / Medium / Low — chọn từ dropdown"),
    ("Environment", 22, "run",
     "Trình duyệt, OS, thiết bị, bản build. Ví dụ: Chrome 115, Windows 11"),
    ("Created By / Date", 20, "gen",
     "Người tạo testcase và ngày tạo"),
    ("Executed By / Date", 20, "run",
     "Người chạy test và ngày chạy"),
]

STATUS_VALUES = ["Pass", "Fail", "Blocked", "Not Executed"]
PRIORITY_VALUES = ["High", "Medium", "Low"]

# Ví dụ thật lấy từ pilot login-sso (SRS §3.2.1 F-01) — không phải ví dụ chế
EXAMPLE = {
    "Test Case ID": "TC_AUTH_001",
    "Module / Feature Name": "User Authentication - Login",
    "Test Title / Summary": "GET /auth/config trả đủ 4 field cấu hình Keycloak",
    "Description / Objective": "Xác nhận FE lấy được cấu hình Keycloak để khởi tạo adapter OIDC",
    "Preconditions": "Portal API sống tại dev-qc-portal.datafactory.fit",
    "Test Data": "(không cần dữ liệu vào)",
    "Test Steps": "1. Gọi GET /api/v1/auth/config, không kèm token\n"
                  "2. Kiểm tra HTTP status\n"
                  "3. Kiểm tra đủ 4 field và đúng kiểu dữ liệu\n"
                  "4. Kiểm tra giá trị keycloak_realm",
    "Expected Result": "1. HTTP 200\n"
                       "2. Có đủ keycloak_enabled, keycloak_issuer, keycloak_client_id, keycloak_realm\n"
                       "3. keycloak_enabled là boolean\n"
                       "4. keycloak_realm = vd-dev",
    "Actual Result": "(tool điền khi chạy)",
    "Status": None,
    "Priority / Severity": "High",
    "Environment": "(tool điền khi chạy)",
    "Created By / Date": "(tool điền khi sinh)",
    "Executed By / Date": "(tool điền khi chạy)",
}


def build_main_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Test Cases"

    for idx, (name, width, group, hint) in enumerate(COLUMNS, start=1):
        col = get_column_letter(idx)
        ws.column_dimensions[col].width = width

        cell = ws.cell(row=1, column=idx, value=name)
        cell.fill = GEN_HDR_FILL if group == "gen" else RUN_HDR_FILL
        cell.font = GEN_HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        cell.comment = Comment(hint, "qc-tool", height=110, width=280)

    ws.row_dimensions[1].height = 34

    # Hàng ví dụ
    for idx, (name, *_rest) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=idx, value=EXAMPLE.get(name))
        cell.font = ID_FONT if idx == 1 else BODY_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[2].height = 96

    # Dropdown
    status_dv = DataValidation(
        type="list", formula1=f'"{",".join(STATUS_VALUES)}"',
        allow_blank=True, showDropDown=False,
    )
    status_dv.error = "Chỉ được chọn: " + " / ".join(STATUS_VALUES)
    status_dv.errorTitle = "Giá trị không hợp lệ"
    ws.add_data_validation(status_dv)
    status_dv.add("J2:J2000")

    priority_dv = DataValidation(
        type="list", formula1=f'"{",".join(PRIORITY_VALUES)}"',
        allow_blank=True, showDropDown=False,
    )
    priority_dv.error = "Chỉ được chọn: " + " / ".join(PRIORITY_VALUES)
    priority_dv.errorTitle = "Giá trị không hợp lệ"
    ws.add_data_validation(priority_dv)
    priority_dv.add("K2:K2000")

    # Tô màu theo Status — nhìn là biết ngay case nào lỗi
    rng = "J2:J2000"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pass"'], fill=GREEN))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fail"'], fill=RED))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Blocked"'], fill=AMBER))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Not Executed"'], fill=GREY))

    ws.freeze_panes = "C2"          # giữ ID + Module khi cuộn ngang
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}2000"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "1F4E79"


def build_guide_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Hướng dẫn")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 96
    ws.sheet_view.showGridLines = False

    rows = [
        ("h1", "HƯỚNG DẪN DÙNG SHEET TESTCASE", ""),
        ("p", "", "Sheet 'Test Cases' có 14 cột. Một dòng = một testcase."),
        ("p", "", "Các bước nằm chung trong ô 'Test Steps', đánh số 1. 2. 3."),
        ("blank", "", ""),
        ("h2", "MÀU HEADER NÓI LÊN ĐIỀU GÌ", ""),
        ("p", "", "Xanh đậm  = tool tự điền khi sinh testcase. Anh review rồi sửa nếu cần."),
        ("p", "", "Nâu cam   = để trống lúc sinh, tool điền vào khi CHẠY test."),
        ("blank", "", ""),
        ("h2", "NHÓM 1 — TOOL TỰ ĐIỀN", ""),
    ]
    for name, _w, group, hint in COLUMNS:
        if group == "gen":
            rows.append(("col", name, hint))
    rows += [
        ("blank", "", ""),
        ("h2", "NHÓM 2 — TOOL ĐIỀN KHI CHẠY", ""),
    ]
    for name, _w, group, hint in COLUMNS:
        if group == "run":
            rows.append(("col", name, hint))
    rows += [
        ("blank", "", ""),
        ("h2", "GIÁ TRỊ HỢP LỆ", ""),
        ("col", "Status", " / ".join(STATUS_VALUES)),
        ("col", "Priority / Severity", " / ".join(PRIORITY_VALUES)),
        ("blank", "", ""),
        ("h2", "CÁCH DÙNG", ""),
        ("p", "", "1. Nhận file từ tool → review nhóm 1, sửa nếu cần."),
        ("p", "", "2. Chạy test → nhóm 2 tự điền."),
        ("p", "", "3. Lọc cột 'Status' = Fail để xem các case lỗi."),
        ("p", "", "4. Hàng số 2 là ví dụ thật (lấy từ login-sso) — xoá được."),
    ]

    r = 1
    for kind, a, b in rows:
        if kind == "blank":
            r += 1
            continue
        ca, cb = ws.cell(row=r, column=1, value=a or None), ws.cell(row=r, column=2, value=b or None)
        if kind == "h1":
            ca.font = Font(size=13, bold=True, color="1F4E79")
        elif kind == "h2":
            ca.font = Font(size=10, bold=True, color="9C5700")
        elif kind == "col":
            ca.font = Font(size=10, bold=True)
            cb.font = BODY_FONT
            cb.alignment = Alignment(wrap_text=True, vertical="top")
        else:
            cb.font = BODY_FONT
        r += 1


def main() -> None:
    wb = Workbook()
    build_main_sheet(wb)
    build_guide_sheet(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Đã ghi {OUT.relative_to(REPO)}  ({OUT.stat().st_size:,} bytes)")
    print(f"  {len(COLUMNS)} cột · {sum(1 for c in COLUMNS if c[2] == 'gen')} cột tool sinh "
          f"· {sum(1 for c in COLUMNS if c[2] == 'run')} cột điền khi chạy")


if __name__ == "__main__":
    main()
