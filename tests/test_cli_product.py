"""Product CLI uses real generation/render/runner; browser/provider stubbed offline."""
from dataclasses import replace
from pathlib import Path
import json

import pytest
import yaml

from qctool.cli import main, build_parser
from qctool.cli_run import build_summary
import qctool.pipeline as pipeline
from qctool.data.fixtures import FixtureStore
from qctool.ui.crawl import UIMap, PageMap
from tests.test_pipeline_generation import setup, route_case, sources, BASE


class Browser:
    evidence = type('Evidence',(),{'to_dict':lambda self:{}})()
    authenticated = False
    def __init__(self,*args,**kwargs): pass
    def start(self): pass
    def close(self): pass
    def start_case(self,*args): pass
    def end_case(self,*args): pass
    def goto(self,url): return BASE+url
    def url_matches(self,pattern): return True,'Exact route observed'
    def capture_failure(self): pass


@pytest.fixture(autouse=True)
def no_live_provider(monkeypatch):
    import qctool.config as config
    monkeypatch.setattr(config,'load_dotenv',lambda *a:None)
    monkeypatch.delenv('QCTOOL_LLM_ENABLED',raising=False)
    monkeypatch.setattr(pipeline,'load_llm_provider',lambda:pytest.fail('Unexpected provider config'))
    import qctool.ui as ui
    import qctool.execute.browser as browser
    monkeypatch.setattr(browser,'BrowserSession',Browser)
    monkeypatch.setattr(ui,'crawl',lambda *a,**kw:UIMap(BASE,'recorded',[PageMap(BASE+'/login','Login')]))


def inputs(monkeypatch,tmp_path, *, source=None, skip=False, cases=None, options=None):
    cfg = setup(monkeypatch,tmp_path,data=source,cases=(route_case(),) if cases is None else cases,web_url=BASE)
    doc = {'docs_dir':str(cfg.docs_dir),'web_url':BASE,'run':{'base_url':BASE},
           'auth':{'mode':'manual'},'out':{'dir':str(cfg.out_dir),'skip_run':skip}}
    doc.update(options or {})
    path=tmp_path/'qctool.yaml';path.write_text(yaml.safe_dump(doc))
    return path,cfg.out_dir


def test_successful_full_run_summary_agrees_with_every_artifact(monkeypatch,tmp_path,capsys):
    path,out=inputs(monkeypatch,tmp_path)
    assert main(['run','--config',str(path)])==0
    summary=json.loads((out/'summary.json').read_text())
    generation=json.loads((out/'3-generate/generation.json').read_text())
    result=json.loads((out/'5-execute/results.json').read_text())
    assert summary['generation_success']==generation['success'] is True
    assert summary['counts']['Pass']==len([r for r in result['results'] if r['status']=='Pass'])==1
    assert summary['executable_testcases']==len(generation['executable_cases'])==1
    assert summary['generated_testcases']==len(yaml.safe_load((out/'3-generate/testcases.yaml').read_text())['testcases'])
    assert summary['executable_coverage']['covered']==summary['executable_coverage']['total']==1
    import csv
    import xml.etree.ElementTree as ET
    from openpyxl import load_workbook
    run_dir=out/'5-execute'
    csv_rows=list(csv.DictReader((run_dir/'results.csv').open(encoding='utf-8-sig')))
    assert [r['Kết quả'] for r in csv_rows]==['Pass']
    junit=ET.parse(run_dir/'junit.xml').getroot()
    assert junit.attrib['tests']=='1' and junit.attrib['failures']=='0'
    workbook=load_workbook(out/'4-render/testcases.xlsx')
    assert any('Pass' in row for row in workbook.active.iter_rows(values_only=True))
    workbook.close()
    for key in ['1-ingest/requirements.yaml','3-generate/testcases.yaml','4-render/testcases.xlsx',
                '3-generate/generation.json','summary.json','pipeline.log','run.json','events.jsonl',
                '5-execute/results.json','5-execute/results.csv','5-execute/junit.xml','6-report/report.md']:
        assert (out/key).is_file(), key
    run=json.loads((out/'run.json').read_text())
    assert run['status']=='done' and run['error'] is None
    assert {k:v['status'] for k,v in run['stages'].items()}=={'ingest':'done','crawl':'done',
        'generate':'done','render':'done','execute':'done','report':'done'}
    assert '6-report/report.md' in run['stages']['report']['outputs']
    for stage in ['1-ingest','3-generate','4-render','5-execute','6-report']:
        assert (out/stage/'stage.log').read_text().strip(), stage
    events=[json.loads(l) for l in (out/'events.jsonl').read_text().splitlines()]
    assert [e['stage'] for e in events][0]=='ingest' and all({'t','stage','level','message'}<=set(e) for e in events)
    text=capsys.readouterr().out
    assert 'QC: SUCCESS' in text and 'Pass: 1' in text
    assert '[generate]' in text and '[ingest]' in text and '[run]' in text and 'Generation:' in text
    assert '[generate]' in (out/'pipeline.log').read_text()


def test_minimal_docs_url_no_llm_no_provisioning(monkeypatch,tmp_path,capsys):
    path,out=inputs(monkeypatch,tmp_path)
    original=pipeline.run_generation;seen=[]
    def spy(*a,**kw):seen.append(kw);return original(*a,**kw)
    monkeypatch.setattr(pipeline,'run_generation',spy)
    cfg=yaml.safe_load(path.read_text())
    assert main(['run','--docs',cfg['docs_dir'],'--url',BASE,'--out',str(out)])==0
    assert seen[0]['resolution_llm'] is None
    assert seen[0]['allow_provisioning'] is False
    assert seen[0]['environment'] is None


def test_llm_enabled_factory_passes_shared_stub_and_executable_case_skips_calls(monkeypatch,tmp_path):
    class Stub:
        def complete(self,prompt):pytest.fail('Deterministic success must skip model')
    stub=Stub();factories=[];seen=[]
    path,_=inputs(monkeypatch,tmp_path,options={'llm':{'enabled':True}})
    monkeypatch.setattr(pipeline,'load_llm_provider',lambda:factories.append(1) or stub)
    original=pipeline.run_generation
    def spy(*a,**kw):seen.append(kw);return original(*a,**kw)
    monkeypatch.setattr(pipeline,'run_generation',spy)
    assert main(['run','--config',str(path)])==0
    assert factories==[1]
    assert seen[0]['resolution_llm'] is stub


def test_provisioning_explicit_gate_forwards_existing_engine_policy(monkeypatch,tmp_path):
    path,_=inputs(monkeypatch,tmp_path,options={'provisioning':{'enabled':True,'environment':'test'}})
    original=pipeline.run_generation;seen=[]
    def spy(*a,**kw):seen.append(kw);return original(*a,**kw)
    monkeypatch.setattr(pipeline,'run_generation',spy)
    assert main(['run','--config',str(path)])==0
    assert seen[0]['allow_provisioning'] and seen[0]['environment']=='test'


def test_unresolved_source_remains_not_executed_and_generation_exit_distinct(monkeypatch,tmp_path,capsys):
    path,out=inputs(monkeypatch,tmp_path,source=sources('Create account','Account provisioned'),cases=())
    assert main(['run','--config',str(path)])==3
    data=json.loads((out/'summary.json').read_text())
    assert not data['generation_success'] and data['pipeline_ok']
    assert data['counts']=={'Pass':0,'Fail':0,'Blocked':0,'Not Executed':1}
    assert data['unresolved_cases'][0]['runtime_status']=='Not Executed'
    assert data['user_action_required']
    assert 'Action:' in capsys.readouterr().out


def test_execution_failure_is_not_generation_failure(monkeypatch,tmp_path,capsys):
    path,out=inputs(monkeypatch,tmp_path)
    monkeypatch.setattr(Browser,'url_matches',lambda *a:(False,'Wrong observed route'))
    assert main(['run','--config',str(path)])==1
    data=json.loads((out/'summary.json').read_text())
    assert data['generation_success'] and data['pipeline_ok']
    assert data['counts']['Fail']==1 and data['status']=='EXECUTION_FAILED'
    assert 'assertion' in capsys.readouterr().out


def test_missing_browser_is_execution_incomplete_not_product_fail(monkeypatch,tmp_path):
    path,out=inputs(monkeypatch,tmp_path)
    monkeypatch.setattr(Browser,'start',lambda *a:(_ for _ in ()).throw(RuntimeError('Browser unavailable')))
    assert main(['run','--config',str(path)])==4
    data=json.loads((out/'summary.json').read_text())
    assert data['counts']['Blocked']==1 and data['counts']['Fail']==0
    assert data['generation_success']


def test_skipped_run_never_claims_pass(monkeypatch,tmp_path):
    path,out=inputs(monkeypatch,tmp_path,skip=True)
    assert main(['run','--config',str(path)])==4
    data=json.loads((out/'summary.json').read_text())
    assert not data['execution_performed'] and data['counts']['Pass']==0
    assert data['counts']['Not Executed']==1


def test_error_secret_redaction_in_terminal_and_log(monkeypatch,tmp_path,capsys):
    secret='PRIVATE-PRODUCT-KEY-DO-NOT-PRINT'
    path,out=inputs(monkeypatch,tmp_path)
    monkeypatch.setenv('QCTOOL_LLM_API_KEY',secret)
    original=pipeline.run_pipeline
    def broken(*a,**kw):
        result=original(*a,**kw);result.ok=False;result.error=secret;result.log.append('Provider '+secret)
        print('stdout '+secret);return result
    monkeypatch.setattr(pipeline,'run_pipeline',broken)
    assert main(['run','--config',str(path)])==2
    assert secret not in capsys.readouterr().out
    assert secret not in (out/'summary.json').read_text()
    assert secret not in (out/'pipeline.log').read_text()
    assert '[REDACTED]' in (out/'pipeline.log').read_text()


def test_progress_is_visible_and_flushed_to_log_before_pipeline_finishes(monkeypatch,tmp_path,capsys):
    secret='PRIVATE-LIVE-PROVIDER-KEY'
    password='PRIVATE-LIVE-PASSWORD'
    path,out=inputs(monkeypatch,tmp_path,options={'auth':{'mode':'manual','password':password}})
    monkeypatch.setenv('QCTOOL_LLM_API_KEY',secret)
    original=pipeline.run_pipeline
    seen=[]
    def spy(cfg,on_progress=None,**kwargs):
        assert callable(on_progress)
        on_progress('crawl','Opening browser '+secret+' '+password)
        visible=capsys.readouterr().out
        persisted=(out/'pipeline.log').read_text()
        assert '[crawl] Opening browser' in visible and '[crawl] Opening browser' in persisted
        assert '[REDACTED]' in visible and '[REDACTED]' in persisted
        assert secret not in visible+persisted and password not in visible+persisted
        assert 'QC: SUCCESS' not in visible  # Callback runs before final summary.
        print('Internal dependency '+secret+' '+password)
        assert capsys.readouterr().out == ''  # Raw stdout remains captured.
        seen.append(True)
        return original(cfg,on_progress=on_progress,**kwargs)
    monkeypatch.setattr(pipeline,'run_pipeline',spy)
    assert main(['run','--config',str(path)]) == 0 and seen == [True]
    log=(out/'pipeline.log').read_text()
    assert '[crawl] Opening browser' in log and 'Internal dependency [REDACTED]' in log
    assert secret not in log and password not in log
    assert 'QC: SUCCESS' in capsys.readouterr().out


def test_failed_pipeline_retains_safe_live_progress(monkeypatch,tmp_path,capsys):
    path,out=inputs(monkeypatch,tmp_path)
    def broken(cfg,on_progress=None,**kwargs):
        on_progress('ingest','Reading source documents')
        raise RuntimeError('Private internal details')
    monkeypatch.setattr(pipeline,'run_pipeline',broken)
    assert main(['run','--config',str(path)]) == 2
    text=capsys.readouterr().out
    assert '[ingest] Reading source documents' in text and 'LỖI PIPELINE' in text
    assert 'Private internal details' not in text
    assert '[ingest] Reading source documents' in (out/'pipeline.log').read_text()


@pytest.mark.parametrize('flags', [[],['--url','javascript:alert(1)'],['--config','/missing/config.yaml']])
def test_preflight_errors_are_controlled(flags,capsys):
    assert main(['run',*flags])==2
    assert 'PREFLIGHT' in capsys.readouterr().out


def test_preflight_only_no_execution_or_artifact_write(monkeypatch,tmp_path,capsys):
    path,out=inputs(monkeypatch,tmp_path)
    monkeypatch.setattr(pipeline,'run_pipeline',lambda *a,**kw:pytest.fail('Preflight must not run'))
    assert main(['check','--config',str(path)])==0
    assert not out.exists() and 'PREFLIGHT OK' in capsys.readouterr().out


def test_bad_yaml_error_never_echoes_secret(monkeypatch,tmp_path,capsys):
    path=tmp_path/'bad.yaml';path.write_text('password: [PRIVATE-SECRET:')
    assert main(['check','--config',str(path)])==2
    assert 'PRIVATE-SECRET' not in capsys.readouterr().out


def test_existing_testcase_yaml_runs_via_step_exec(monkeypatch,tmp_path):
    import qctool.cli as cli
    seen=[]
    monkeypatch.setattr(cli,'cmd_run_testcases',lambda args:seen.append(args.testcases) or 0)
    assert cli.main(['step','exec','old.yaml','--base-url',BASE])==0 and seen==['old.yaml']


def test_init_creates_project_and_features_and_targets_resolve(monkeypatch,tmp_path):
    import qctool.cli as cli
    monkeypatch.chdir(tmp_path)
    assert cli.main(['init','shop'])==0 and (tmp_path/'projects/shop/project.yaml').is_file()
    assert cli.main(['init','shop'])==2
    assert cli.main(['init','shop/login'])==0 and cli.main(['init','shop/cart'])==0
    assert cli.main(['init','shop/cart'])==2
    feats=lambda *a: [p.parent.name for p in cli.resolve_targets(*a)]
    assert feats('shop')==['cart','login']
    assert feats('shop',['login'])==['login']==feats('projects/shop/login')==feats('shop/login')
    with pytest.raises(SystemExit, match='nope'):
        cli.resolve_targets('nope')
    with pytest.raises(SystemExit, match='không có feature'):
        cli.resolve_targets('shop',['nope'])
    with pytest.raises(SystemExit, match='--cart'):
        cli.resolve_project('shop')


def test_run_project_runs_every_feature_or_only_flagged_one(monkeypatch,tmp_path,capsys):
    import qctool.cli as cli
    import qctool.cli_run as cli_run
    monkeypatch.chdir(tmp_path)
    cli.main(['init','shop/login']); cli.main(['init','shop/cart'])
    seen=[]
    monkeypatch.setattr(cli_run,'cmd_project',lambda args,validate_only=False:seen.append(Path(args.config).parent.name) or (1 if 'cart' in args.config else 0))
    assert cli.main(['run','shop'])==1 and seen==['cart','login']
    assert 'Tổng hợp' in capsys.readouterr().out
    seen.clear(); assert cli.main(['run','shop','--login'])==0 and seen==['login']
    seen.clear(); assert cli.main(['check','shop','-f','cart'])==1 and seen==['cart']
    with pytest.raises(SystemExit):
        cli.main(['run','shop','--nope'])


def test_missing_locked_fixture_shows_exact_need_without_verifying_or_provisioning(monkeypatch,tmp_path,capsys):
    options={'test_fixtures':{
        'fixtures':{'locked_user':{'email':{'env':'MISSING_EMAIL'},'password':{'env':'MISSING_PASSWORD'}}},
        'bindings':{'FLOW-1':{'email':'locked_user.email','password':'locked_user.password'}},
        'acquisition':{'requirements':{'FLOW-1':[{'fixture':'locked_user',
            'aliases':['locked_user.email','locked_user.password'],
            'facts':{'account_exists':True,'role':'customer','account_state':'locked'}}]}}}}
    monkeypatch.delenv('MISSING_EMAIL',raising=False);monkeypatch.delenv('MISSING_PASSWORD',raising=False)
    path,out=inputs(monkeypatch,tmp_path,options=options,cases=())
    assert main(['run','--config',str(path)])==3
    summary=json.loads((out/'summary.json').read_text())
    assert summary['counts']['Pass']==0 and summary['counts']['Not Executed']==summary['generated_testcases']
    assert summary['unresolved_test_data'][0]['code']=='MISSING_FIXTURE'
    assert summary['unresolved_test_data'][0]['needs'][0]['facts']['account_state']=='locked'
    text=capsys.readouterr().out
    assert 'locked_user' in text and 'locked' in text and 'MISSING_FIXTURE' in text


def test_unassessable_requirement_prevents_false_zero_of_zero_success(monkeypatch,tmp_path):
    source={'requirements':[{'id':'FR-1','kind':'functional','title':'Unknown behavior'}], 'counts':{'total':1}}
    path,out=inputs(monkeypatch,tmp_path,source=source)
    assert main(['run','--config',str(path)])==3
    coverage=json.loads((out/'summary.json').read_text())['executable_coverage']
    assert coverage['unassessable_sources']==1 and not coverage['complete']


def test_openapi_preflight_reuses_loader_without_network(monkeypatch,tmp_path,capsys):
    spec=tmp_path/'api.yaml'
    spec.write_text(yaml.safe_dump({'openapi':'3.0.0','info':{'title':'Example','version':'1'},
        'paths':{'/health':{'get':{'responses':{'200':{'description':'OK'}}}}}}))
    assert main(['check','--openapi',str(spec),'--url',BASE])==0
    assert 'API operations: 1' in capsys.readouterr().out
    spec.write_text('paths: []')
    assert main(['check','--openapi',str(spec),'--url',BASE])==2


def test_inline_secret_is_redacted_using_config_not_just_environment(monkeypatch,tmp_path,capsys):
    secret='INLINE-SECRET-NOT-A-REAL-PASSWORD'
    path,out=inputs(monkeypatch,tmp_path,options={'auth':{'mode':'manual','password':secret}})
    original=pipeline.run_pipeline
    def spy(*a,**kw):
        result=original(*a,**kw);result.log.append(secret);print(secret);return result
    monkeypatch.setattr(pipeline,'run_pipeline',spy)
    assert main(['run','--config',str(path)])==0
    assert secret not in capsys.readouterr().out and secret not in (out/'pipeline.log').read_text()


def test_provisioning_boolean_and_environment_are_explicit(monkeypatch,tmp_path,capsys):
    path,_=inputs(monkeypatch,tmp_path,options={'provisioning':{'enabled':'false'}})
    assert main(['run','--config',str(path)])==2
    assert 'boolean' in capsys.readouterr().out
    doc=yaml.safe_load(path.read_text());doc['provisioning']={'enabled':True}
    path.write_text(yaml.safe_dump(doc))
    assert main(['run','--config',str(path)])==2
    assert 'environment' in capsys.readouterr().out


def test_cli_overrides_config_disable_llm_and_provisioning(monkeypatch,tmp_path):
    path,_=inputs(monkeypatch,tmp_path,options={'llm':{'enabled':True},'provisioning':{'enabled':True,'environment':'test'}})
    original=pipeline.run_generation;seen=[]
    def spy(*a,**kw):seen.append(kw);return original(*a,**kw)
    monkeypatch.setattr(pipeline,'run_generation',spy)
    assert main(['run','--config',str(path),'--no-llm','--no-allow-provisioning'])==0
    assert seen[0]['resolution_llm'] is None and seen[0]['allow_provisioning'] is False
