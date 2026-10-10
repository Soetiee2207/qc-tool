"""Agent safety/schema/citation and repair feedback, no live model."""
from copy import deepcopy
import json
from pathlib import Path
from unittest.mock import Mock
import pytest

from qctool.agent.gate import case_schema, draft_case, check_case
from qctool.agent.safety import SafetyPolicy, SafetyError
from qctool.agent.tools import AgentTools
from qctool.agent.loop import author_intent, AgentBudget
from qctool.agent.knowledge import AppKnowledge
from qctool.testcase import ACTION_SPECS
from qctool.generation import run_generation

INTENT = {'id':'AC-1:then','req_id':'AC-1','given':'anonymous','when':'Open /a','expected':'URL /a','title':'Open A'}

def raw(pattern='/a$'):
    return {'title':'Open A','steps':[
        {'seq':1,'kind':'given','action':'ui.session','args':{'mode':'anonymous'}},
        {'seq':2,'kind':'when','action':'ui.goto','args':{'url':'/a'}},
        {'seq':3,'kind':'then','action':'assert.url_matches','args':{'pattern':pattern},'source':{'req_id':'AC-1','quote':'URL /a'}}]}


def case(value=None):
    return draft_case(value or raw(),case_id='TC_A_001',module='A',intent=INTENT)


def test_schema_derived():
    variants = case_schema()['properties']['steps']['items']['oneOf']
    for v in variants:
        spec = ACTION_SPECS[v['properties']['action']['const']]
        assert v['properties']['args']['required'] == list(spec.required)
        assert set(v['properties']['args']['properties']) == set(spec.required + spec.optional)


def test_draft_exact_schema():
    with pytest.raises(ValueError): draft_case({**raw(),'id':'invented'},case_id='TC_A_001',module='A',intent=INTENT)
    bad=raw();bad['steps'][0]['args']['guess']=True
    with pytest.raises(ValueError): case(bad)


def test_citation_verbatim_and_intent_bound():
    bad=raw();bad['steps'][2]['source']['quote']='URL /b'
    assert not check_case(case(bad),INTENT,SafetyPolicy('http://app.test')).accepted
    bad=raw();bad['steps'][2]['source']['req_id']='AC-2'
    assert not check_case(case(bad),INTENT,SafetyPolicy('http://app.test')).accepted


def test_literal_absent_needs_review_not_ready():
    r=check_case(case(raw('/wrong$')),INTENT,SafetyPolicy('http://app.test'))
    assert r.accepted and r.review


def test_all_outcomes_cited():
    intent={**INTENT,'expected':'URL /a; Heading A visible'}
    r=check_case(case(),intent,SafetyPolicy('http://app.test'))
    assert not r.accepted


@pytest.mark.parametrize('url',['https://evil.test','javascript:alert(1)','file:///tmp/foo','//evil.test','/admin/users'])
def test_safety_targets(url):
    with pytest.raises(SafetyError): SafetyPolicy('http://app.test').url(url)


def test_unreviewed_write_blocked_inside_tool():
    policy=SafetyPolicy('http://app.test')
    route=Mock();route.request.url='http://app.test/api/orders';route.request.method='POST'
    policy._route(route)
    route.abort.assert_called_once();route.continue_.assert_not_called()


def test_foreign_read_blocked_inside_tool():
    policy=SafetyPolicy('http://app.test')
    route=Mock();route.request.url='http://evil.test/a';route.request.method='GET'
    policy._route(route)
    route.abort.assert_called_once()


def test_dry_run_never_calls_assertion_or_returns_values(monkeypatch):
    tools=AgentTools(INTENT,'TC_A_001','A',SafetyPolicy('http://app.test'))
    tools.candidate=case()
    browser=Mock()
    monkeypatch.setattr(tools,'_open',lambda:browser)
    monkeypatch.setattr(tools,'act',lambda *a,**k:{'ok':True})
    result=tools.dry_run()
    browser.url_matches.assert_not_called()
    assert result['ok'] and 'actual' not in json.dumps(result)
    assert tools.ready.generation_status == 'ready'


class Replay:
    def __init__(self, rounds): self.rounds=list(rounds);self.calls=[]
    def chat(self,messages,tools,**kwargs):
        self.calls.append(deepcopy(messages))
        name,args=self.rounds.pop(0)
        return {'role':'assistant','content':None,'tool_calls':[{'id':str(len(self.calls)), 'type':'function','function':{'name':name,'arguments':json.dumps(args)}}]}


def test_error_feedback_same_session_and_no_input_mutation(monkeypatch):
    source=deepcopy(INTENT)
    tools=AgentTools(source,'TC_A_001','A',SafetyPolicy('http://app.test'))
    monkeypatch.setattr(tools,'_open',lambda:Mock())
    monkeypatch.setattr(tools,'act',lambda *a,**k:{'ok':True})
    llm=Replay([('draft_case',{'title':'bad'}),('draft_case',raw()),('dry_run',{})])
    result=author_intent(source,tools,llm,budget=AgentBudget(max_tokens=100000))
    assert result.dry_run_passed and result.case.generation_status=='ready'
    assert llm.calls[1][-1]['role']=='tool' and 'error' in llm.calls[1][-1]['content']
    assert source==INTENT


def test_repeated_error_stops_without_pingpong():
    tools=AgentTools(INTENT,'TC_A_001','A',SafetyPolicy('http://app.test'))
    llm=Replay([('draft_case',{}),('draft_case',{})])
    result=author_intent(INTENT,tools,llm,budget=AgentBudget(max_tokens=100000))
    assert result.stop_reason=='REPEATED_ERROR' and len(llm.calls)==2


def test_provider_failure_preserves_manual_source():
    tools=AgentTools(INTENT,'TC_A_001','A',SafetyPolicy('http://app.test'))
    llm=Mock();llm.chat.side_effect=TimeoutError('SECRET')
    result=author_intent(INTENT,tools,llm)
    assert result.case.source_quote==INTENT['expected']
    assert 'SECRET' not in json.dumps(result.to_dict())


def test_no_llm_stays_manual_and_immutable():
    rows=[{'id':'AC-1','kind':'acceptance_criterion','given':'anonymous','when':'Open /a','then':'URL /a'}]
    before=deepcopy(rows)
    result=run_generation(rows,module='A')
    assert rows==before and result.unresolved_cases and not result.success
    assert result.cases[0].generation_status=='manual'


def test_knowledge_no_auth_or_raw_secrets(tmp_path):
    from qctool.test_fixtures import FixtureStore
    store=FixtureStore({'user':{'password':{'env':'PW'}}},environ={'PW':'SECRET_PASSWORD'})
    kb=AppKnowledge(tmp_path/'app-model.yaml','http://app.test',store=store)
    kb.observe({'url':'http://app.test/a','aria':'SECRET_PASSWORD','stable_locators':[]})
    kb.save()
    assert 'SECRET_PASSWORD' not in kb.path.read_text()
    with pytest.raises(ValueError): kb.remember('bad',[{'action':'ui.fill','args':{'target':'pw','value':'SECRET_PASSWORD'}}])


def test_design_cannot_drop_or_invent_source():
    from qctool.agent.loop import design_intents
    with pytest.raises(ValueError):
        design_intents([INTENT],Replay([('design_tests',{'intents':[]})]))
    rows,_=design_intents([INTENT],Replay([('design_tests',{'intents':[{'id':INTENT['id'],'title':'A','source_quote':INTENT['expected']}]})]))
    assert rows[0]['expected']==INTENT['expected'] and rows[0]['when']==INTENT['when']


def test_runtime_does_not_infer_expected_from_dry_run(monkeypatch):
    tools=AgentTools(INTENT,'TC_A_001','A',SafetyPolicy('http://app.test'))
    tools.candidate=case(raw('/wrong$'))
    monkeypatch.setattr(tools,'_open',lambda:Mock())
    monkeypatch.setattr(tools,'act',lambda *a,**k:{'ok':True})
    assert tools.dry_run()['ok']
    assert tools.ready.generation_status=='needs_review'
    assert tools.ready.steps[-1].args['pattern']=='/wrong$'


def test_stable_prefix_and_step_order():
    from qctool.agent.loop import protocol
    assert protocol()==protocol()
    assert [s.seq for s in case().steps]==[1,2,3]


def test_fixture_alias_requires_current_source_binding():
    from qctool.test_fixtures import FixtureStore
    store=FixtureStore({'user':{'password':{'env':'PW'}}},bindings={'OTHER':{'password':'user.password'}},environ={'PW':'SECRET'})
    policy=SafetyPolicy('http://app.test',fixture_store=store,req_ids=['AC-1'])
    with pytest.raises(SafetyError):policy.value('user.password','source')


def test_macro_stale_locator_feedback(tmp_path,monkeypatch):
    kb=AppKnowledge(tmp_path/'app-model.yaml','http://app.test')
    kb.remember('navigate',[{'action':'ui.goto','args':{'url':'/a'}}])
    tools=AgentTools(INTENT,'TC_A_001','A',SafetyPolicy('http://app.test'),knowledge=kb)
    monkeypatch.setattr(tools,'act',Mock(side_effect=SafetyError('Stale live state')))
    answer=tools.call('use_macro',{'name':'navigate'})
    assert not answer['ok'] and 'Stale' in answer['error']


def test_list_api_feedback_is_object_not_list():
    tools=AgentTools(INTENT,'TC_A_001','A',SafetyPolicy('http://app.test'))
    assert tools.call('list_api',{}) == {'items':[]}
    assert tools.audit[-1]['ok']


def test_unique_stable_id_wins_over_ambiguous_role():
    from qctool.agent.tools import locator
    page=Mock();stable=Mock();stable.count.return_value=1
    page.get_by_test_id.return_value=stable
    assert locator(page,{'target':'submit-id','role':'button','by_name':'Submit'}) is stable
    page.get_by_role.assert_not_called()


def test_duplicate_stable_id_never_uses_name_fallback():
    from qctool.agent.tools import locator
    page=Mock();page.get_by_test_id.return_value.count.return_value=2
    with pytest.raises(SafetyError):locator(page,{'target':'submit-id','role':'button','by_name':'Submit'})
    page.get_by_role.assert_not_called()


def test_read_contract_host_does_not_authorize_browser_document_navigation():
    from types import SimpleNamespace
    ep=SimpleNamespace(method='GET',path='/read',contract_servers=['http://api.test'],servers=[])
    policy=SafetyPolicy('http://app.test',endpoints=[ep])
    route=Mock();route.request.url='http://api.test/read';route.request.method='GET'
    route.request.is_navigation_request.return_value=True
    policy._route(route)
    route.abort.assert_called_once()


def test_review_labels_do_not_change_runtime_results(tmp_path):
    from qctool.render.excel import render
    from qctool.render.excel import read_review_labels
    from openpyxl import load_workbook
    c=case();c.generation_status='needs_review';c.review_reasons=['Weak oracle']
    path=tmp_path/'cases.xlsx';render([c],path)
    wb=load_workbook(path);assert wb['Test Cases'].max_column==14
    wb['Review']['D2']='accepted';wb.save(path);wb.close()
    assert read_review_labels(path)=={c.id:{'review_decision':'accepted'}}
    assert c.status!='Pass'


def test_exact_ascii_accessible_name_without_test_id():
    from qctool.agent.tools import locator
    page=Mock();page.get_by_test_id.return_value.count.return_value=0
    role=page.get_by_role.return_value;role.count.return_value=1
    assert locator(page,{'target':'Email','role':'textbox','by_name':'Email'}) is role
    page.get_by_role.assert_called_once_with('textbox',name='Email',exact=True)


def test_invented_stable_id_cannot_fall_back_to_an_unrelated_name():
    from qctool.agent.tools import locator
    page=Mock();page.get_by_test_id.return_value.count.return_value=0
    page.get_by_test_id.return_value.wait_for.side_effect=TimeoutError()
    with pytest.raises(TimeoutError):locator(page,{'target':'invented-submit','role':'button','by_name':'Submit'})
    page.get_by_role.assert_not_called()


def test_empty_expected_value_is_a_valid_oracle_not_missing():
    from qctool.testcase import Step, StepKind, ValidationError
    Step(1, StepKind.THEN, 'assert.value', {'target': 'email', 'expected': ''}).validate()
    with pytest.raises(ValidationError):
        Step(1, StepKind.THEN, 'assert.value', {'target': '', 'expected': 'x'}).validate()


def test_capture_error_explains_shape():
    from qctool.testcase import Step, StepKind, ValidationError
    with pytest.raises(ValidationError) as err:
        Step(1, StepKind.WHEN, 'ui.click', {'target': 'b', 'capture_response': {'method': 'POST', 'path': '/x'}}).validate()
    assert 'method/path/origin' in str(err.value) and 'origin' in str(err.value)


def test_protocol_documents_capture_contract():
    from qctool.agent.loop import protocol
    text = protocol()
    assert 'capture_requests' in text and '$.request_count' in text and 'expected ""' in text


def test_empty_expected_value_is_executable():
    from qctool.executability import analyze_executability
    from qctool.testcase import Step, StepKind, TestCase
    c = TestCase('TC_E_001', 'E', 'Empty', 'Email trống', traces_to=['AC-1'], source_quote='Email trống',
                 steps=[Step(1, StepKind.THEN, 'assert.value', {'target': 'email', 'expected': ''})])
    assert not [g for g in analyze_executability(c).gaps if g.code == 'MISSING_ARGUMENT']


def test_high_risk_route_opens_only_when_source_names_it():
    from qctool.agent.gate import source_routes
    routes = source_routes({'given': 'customer đã đăng nhập', 'when': 'Mở /admin', 'expected': 'URL chuyển tới /settings/delete'})
    assert routes == ('/admin',)
    allowed = SafetyPolicy('https://app.test', source_routes=routes)
    assert allowed.url('/admin') == 'https://app.test/admin'
    with pytest.raises(SafetyError):
        allowed.url('/admin/delete-user')  # not named by the source
    with pytest.raises(SafetyError):
        SafetyPolicy('https://app.test').url('/admin')


def test_outcome_asserted_before_final_action_is_rejected():
    intent = {**INTENT, 'when': 'Đăng nhập rồi mở /a', 'expected': 'URL /a'}
    early = {'title': 'A', 'steps': [
        {'seq': 1, 'kind': 'given', 'action': 'ui.session', 'args': {'mode': 'anonymous'}},
        {'seq': 2, 'kind': 'then', 'action': 'assert.url_matches', 'args': {'pattern': '/a$'}, 'source': {'req_id': 'AC-1', 'quote': 'URL /a'}},
        {'seq': 3, 'kind': 'when', 'action': 'ui.goto', 'args': {'url': '/a'}}]}
    c = draft_case(early, case_id='TC_A_001', module='A', intent=intent)
    assert any('after the final action' in e for e in check_case(c, intent, SafetyPolicy('https://app.test')).errors)
    assert check_case(case(), INTENT, SafetyPolicy('https://app.test')).accepted


def test_source_route_must_be_performed_not_skipped():
    intent = {**INTENT, 'when': 'Đăng nhập; mở /admin', 'expected': 'URL /a'}
    skipped = {'title': 'A', 'steps': [
        {'seq': 1, 'kind': 'given', 'action': 'ui.session', 'args': {'mode': 'anonymous'}},
        {'seq': 2, 'kind': 'when', 'action': 'ui.goto', 'args': {'url': '/a'}},
        {'seq': 3, 'kind': 'then', 'action': 'assert.url_matches', 'args': {'pattern': '/a$'}, 'source': {'req_id': 'AC-1', 'quote': 'URL /a'}}]}
    c = draft_case(skipped, case_id='TC_A_001', module='A', intent=intent)
    assert any('/admin' in e for e in check_case(c, intent, SafetyPolicy('https://app.test')).errors)


def test_url_oracle_must_be_a_path_regex_not_prose():
    intent = {**INTENT, 'expected': 'URL chuyển tới /a'}
    def check(pattern):
        value = raw(pattern); value['steps'][2]['source']['quote'] = 'URL chuyển tới /a'
        return check_case(draft_case(value, case_id='TC_A_001', module='A', intent=intent), intent, SafetyPolicy('https://app.test'))
    prose = check('URL chuyển tới /a')
    assert not prose.accepted and any('not prose' in e for e in prose.errors)
    absolute = check('^https://app\\.test/a$')
    assert absolute.accepted and not absolute.review


def test_operator_contract_matches_browser():
    from qctool.testcase import Step, StepKind, ValidationError, allowed_operators
    # browser.text/value run contains/eq/matches via expect(); validator must agree.
    for action in ('assert.text', 'assert.value'):
        assert allowed_operators(action) == {'eq', 'contains', 'matches'}
        Step(1, StepKind.THEN, action, {'target': 't', 'expected': 'x', 'operator': 'contains'}).validate()
        with pytest.raises(ValidationError):
            Step(1, StepKind.THEN, action, {'target': 't', 'expected': 'x', 'operator': 'gt'}).validate()
    # json_field has no `contains` in runner._compare: must stay rejected there.
    with pytest.raises(ValidationError):
        Step(1, StepKind.THEN, 'assert.json_field', {'path': '$.a', 'expected': 1, 'operator': 'contains'}).validate()


def test_protocol_lists_operators_per_verb():
    from qctool.agent.loop import protocol
    vocab = json.loads(protocol().split('Action vocabulary: ', 1)[1])
    assert vocab['assert.text']['operators'] == ['contains', 'eq', 'matches']


def test_timeout_feedback_names_phase_without_raw_text(monkeypatch):
    tools = AgentTools(INTENT, 'TC_A_001', 'A', SafetyPolicy('https://app.test'))
    class TimeoutError(Exception):
        pass
    def boom(*a, **k):
        tools.phase = 'ui.fill login-email-input'
        raise TimeoutError('waiting for https://app.test/x?token=SECRET')
    monkeypatch.setattr(tools, 'act', boom)
    answer = tools.call('browser_act', {'action': 'ui.fill', 'args': {'target': 'login-email-input', 'value': 'x'}})
    assert not answer['ok'] and 'ui.fill login-email-input' in answer['error'] and 'SECRET' not in answer['error']


def test_compact_history_keeps_latest_page_and_pairs():
    from qctool.agent.loop import compact_history
    page = lambda url: json.dumps({'url': url, 'aria': 'X' * 5000, 'stable_locators': [1] * 50})
    messages = [{'role': 'system', 'content': 's'},
                {'role': 'tool', 'tool_call_id': 'a', 'content': page('/one')},
                {'role': 'tool', 'tool_call_id': 'b', 'content': json.dumps({'ok': False, 'error': 'keep me'})},
                {'role': 'tool', 'tool_call_id': 'c', 'content': page('/two')}]
    compact_history(messages)
    old, kept = json.loads(messages[1]['content']), json.loads(messages[3]['content'])
    assert 'aria' not in old and old['url'] == '/one' and old['superseded']
    assert kept['aria'] and messages[1]['tool_call_id'] == 'a'
    assert json.loads(messages[2]['content'])['error'] == 'keep me'


def test_protocol_warns_capture_window_hides_transient_messages():
    from qctool.agent.loop import protocol
    assert '"timeout_ms": 2000' in protocol()
