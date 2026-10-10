"""Test cho chặng execute + report.

Dùng mock server thật (dựng trong tiến trình, cổng do hệ điều hành chọn) thay vì
giả lập HTTP — để test chạy đúng đường đi thật của runner.
"""

from __future__ import annotations

import sys
import threading
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from qctool.execute import (  # noqa: E402
    STATUS_BLOCKED,
    STATUS_FAIL,
    STATUS_PASS,
    HttpSession,
    run_case,
    run_cases,
)
from qctool.generate import IdAllocator, generate_all  # noqa: E402
from qctool.report import build_markdown, write_csv, write_json, write_junit  # noqa: E402
from qctool.core.testcase import Step, StepKind, TestCase  # noqa: E402
from tests.mock_server import create_server  # noqa: E402

SAMPLE_SPEC = Path(__file__).resolve().parent.parent / "samples" / "openapi" / "task-manager-sample.yaml"


@pytest.fixture(scope="module")
def api():
    server = create_server(0, "/api/v1")
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    host, port = server.server_address[:2]
    yield f"http://{host}:{port}/api/v1"
    server.shutdown()
    server.server_close()


@pytest.fixture
def session(api):
    return HttpSession(api, bearer_token="test-token", timeout=10)


def _case(**kw) -> TestCase:
    base = dict(
        id="TC_AUTH_001", module="AUTH", title="t", objective="o",
        traces_to=["AC-1"], source_quote="q",
    )
    return TestCase(**{**base, **kw})


# ── Chạy case ────────────────────────────────────────────────────────────────
def test_pass_when_expectation_matches(session):
    case = _case(steps=[
        Step(1, StepKind.WHEN, "http.request",
             {"method": "GET", "path": "/health", "auth": "none", "save_as": "res"}),
        Step(2, StepKind.THEN, "assert.http_status", {"expected": 200, "source": "res"}),
    ])
    result = run_case(case, session)
    assert result.status == STATUS_PASS
    assert len(result.steps) == 2
    assert all(s.ok for s in result.steps)


def test_fail_when_expectation_differs(session):
    """Mock trả 200; mong đợi 999 → phải Fail, và ghi rõ thực tế là bao nhiêu."""
    case = _case(steps=[
        Step(1, StepKind.WHEN, "http.request",
             {"method": "GET", "path": "/health", "auth": "none", "save_as": "res"}),
        Step(2, StepKind.THEN, "assert.http_status", {"expected": 999, "source": "res"}),
    ])
    result = run_case(case, session)
    assert result.status == STATUS_FAIL
    assert any("999" in d and "200" in d for d in result.actual)


def test_assert_stops_at_first_failure(session):
    """Fail ở bước kiểm tra đầu thì không chạy tiếp — tránh báo lỗi dây chuyền."""
    case = _case(steps=[
        Step(1, StepKind.WHEN, "http.request",
             {"method": "GET", "path": "/health", "auth": "none", "save_as": "res"}),
        Step(2, StepKind.THEN, "assert.http_status", {"expected": 999, "source": "res"}),
        Step(3, StepKind.THEN, "assert.json_schema",
             {"required": ["status"], "source": "res"}),
    ])
    result = run_case(case, session)
    assert result.status == STATUS_FAIL
    assert len(result.steps) == 2, "phải dừng ở bước lỗi đầu tiên"


def test_unmapped_case_is_blocked_not_failed(session):
    """Case còn bước văn xuôi → Blocked, KHÔNG phải Fail.

    Đây là điểm quan trọng: báo Fail nghĩa là sản phẩm có lỗi. Case chưa map thì
    chưa kiểm tra được gì cả, nên phải là Blocked.
    """
    case = _case(steps=[
        Step(1, StepKind.WHEN, "perform.unmapped", {"description": "bấm nút"}),
        Step(2, StepKind.THEN, "assert.unmapped", {"expected": "thấy trang chủ"}),
    ])
    result = run_case(case, session)
    assert result.status == STATUS_BLOCKED
    assert "map" in (result.blocked_reason or "").lower()


def test_case_needing_setup_is_blocked_not_failed(session):
    """Case cần môi trường (tài khoản thiếu quyền…) → Blocked, không Failed."""
    case = _case(requires="Cần một tài khoản KHÔNG có quyền",
                 steps=[Step(1, StepKind.THEN, "assert.http_status", {"expected": 403})])
    result = run_case(case, session)
    assert result.status == STATUS_BLOCKED
    assert result.blocked_reason == "Cần một tài khoản KHÔNG có quyền"


def test_connection_error_reports_clearly():
    """Không kết nối được phải báo rõ, không sập runner."""
    dead = HttpSession("http://127.0.0.1:1/api", bearer_token="x", timeout=2)
    case = _case(steps=[
        Step(1, StepKind.WHEN, "http.request",
             {"method": "GET", "path": "/health", "auth": "none", "save_as": "res"}),
        Step(2, StepKind.THEN, "assert.http_status", {"expected": 200, "source": "res"}),
    ])
    result = run_case(case, dead)
    assert result.status == STATUS_FAIL
    assert result.blocked_reason


def test_path_params_are_substituted(session):
    """Bug đã gặp: `{project_id}` không được thay → mọi case có path param nhận 404.

    Đây là test hồi quy cho lỗi đó.
    """
    endpoints, _, _ = __import__(
        "qctool.generate", fromlist=["load_endpoints"]
    ).load_endpoints(SAMPLE_SPEC)
    project_id = "3f2504e0-4f89-11d3-9a0c-0305e82c3301"
    alloc = IdAllocator()
    from qctool.generate.from_openapi import generate_for_endpoint

    get_project = next(e for e in endpoints
                       if e.path == "/projects/{project_id}" and e.method == "GET")
    cases = generate_for_endpoint(get_project, alloc)
    happy = next(c for c in cases if c.case_type.value == "positive")

    call = happy.steps[0]
    assert "{" not in call.args["path"], f"path chưa được thay: {call.args['path']}"
    assert project_id in call.args["path"]

    result = run_case(happy, session)
    assert result.status == STATUS_PASS, result.actual


def test_malformed_json_body_sent_raw(session):
    """Case 400 phải gửi JSON hỏng thật, không phải JSON hợp lệ chứa field lạ."""
    case = _case(steps=[
        Step(1, StepKind.WHEN, "http.request", {
            "method": "POST", "path": "/projects", "auth": "bearer",
            "raw_body": '{"name": "abc", ', "save_as": "res",
        }),
        Step(2, StepKind.THEN, "assert.http_status", {"expected": 400, "source": "res"}),
    ])
    result = run_case(case, session)
    assert result.status == STATUS_PASS, result.actual


# ── Chạy cả lô ───────────────────────────────────────────────────────────────
def test_run_all_generated_cases(session):
    cases, _, _ = generate_all(openapi_path=SAMPLE_SPEC, module_override="SAMPLE")
    result = run_cases(cases, session, environment="mock")
    assert result.total == len(cases)
    assert sum(result.counts.values()) == result.total
    # Mock có bug cài sẵn → phải có Fail. Nếu không có Fail nào thì runner đang
    # không thật sự kiểm tra gì cả.
    assert result.counts[STATUS_FAIL] > 0


def test_planted_bugs_are_detected(session):
    """Ba bug cài sẵn trong mock phải bị bắt.

    Tìm case theo **nội dung kiểm tra** (endpoint + kỳ vọng), không theo ID cứng:
    ID đổi khi tên module hoặc thứ tự sinh thay đổi, nhưng tính chất "case này
    kiểm bug đó" thì không. Test bám ID cứng sẽ vỡ mỗi lần đổi cách đặt tên, dù
    hành vi vẫn đúng.
    """
    cases, _, _ = generate_all(openapi_path=SAMPLE_SPEC, module_override="SAMPLE")
    result = run_cases(cases, session, environment="mock")
    by_id = {c.case_id: c for c in result.cases}

    def find(*, path: str, status: int) -> str:
        for case in result.cases:
            if f"HTTP {status}" not in " ".join(case.actual):
                continue
            if path in case.title or any(path in s.detail for s in case.steps):
                return case.case_id
        raise AssertionError(f"không thấy case nào kiểm {path} → {status}")

    # BUG1: POST /projects trả 200 thay vì 201
    assert by_id[find(path="/projects", status=201)].status == STATUS_FAIL
    # BUG3: /health chậm 127ms > ngưỡng 50ms
    slow = next(c for c in result.cases
                if "VƯỢT ngưỡng" in " ".join(c.actual))
    assert slow.status == STATUS_FAIL
    # BUG4: GET /projects thiếu field page_size
    missing = next(c for c in result.cases
                   if "page_size" in " ".join(c.actual)
                   and c.status == STATUS_FAIL)
    assert missing.status == STATUS_FAIL


def test_report_numbers_add_up(session):
    cases, _, _ = generate_all(openapi_path=SAMPLE_SPEC, module_override="SAMPLE")
    result = run_cases(cases, session, environment="mock")
    counts = result.counts
    assert counts[STATUS_PASS] + counts[STATUS_FAIL] + counts[STATUS_BLOCKED] == result.total
    text = build_markdown(result)
    assert str(result.total) in text
    assert "Tỉ lệ đạt" in text or "Chưa chạy được" in text


# ── Report ghi file ──────────────────────────────────────────────────────────
def test_report_writes_all_formats(tmp_path, session):
    cases, _, _ = generate_all(openapi_path=SAMPLE_SPEC, module_override="SAMPLE")
    result = run_cases(cases, session, environment="mock")

    write_json(result, tmp_path / "results.json")
    write_csv(result, tmp_path / "results.csv")
    write_junit(result, tmp_path / "junit.xml")

    assert (tmp_path / "results.json").is_file()
    assert (tmp_path / "results.csv").is_file()

    xml = (tmp_path / "junit.xml").read_text(encoding="utf-8")
    assert xml.startswith("<?xml")
    assert f'tests="{result.total}"' in xml
    assert f'failures="{result.counts[STATUS_FAIL]}"' in xml

    # CSV dùng BOM để Excel mở đúng tiếng Việt
    raw = (tmp_path / "results.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")


def test_results_written_back_to_xlsx(tmp_path, session):
    """Kết quả phải ghi ngược được vào file template — đây là lý do có 4 cột runtime."""
    import openpyxl

    from qctool.render import render
    from qctool.render.results import fill_results

    cases, _, _ = generate_all(openapi_path=SAMPLE_SPEC, module_override="SAMPLE")
    xlsx = tmp_path / "tc.xlsx"
    render(cases, xlsx)

    # Trước khi chạy: cột Status trống
    ws = openpyxl.load_workbook(xlsx)["Test Cases"]
    assert all(ws.cell(row=r, column=10).value in (None, "") for r in range(2, ws.max_row + 1))

    result = run_cases(cases, session, environment="mock")
    updated = fill_results(xlsx, result)
    assert updated == result.total

    ws = openpyxl.load_workbook(xlsx)["Test Cases"]
    statuses = [ws.cell(row=r, column=10).value for r in range(2, ws.max_row + 1)]
    assert all(s in {STATUS_PASS, STATUS_FAIL, STATUS_BLOCKED} for s in statuses)
    assert any(s == STATUS_FAIL for s in statuses), "mock có bug, phải có Fail"
    # Cột Actual Result cũng phải được điền
    assert ws.cell(row=2, column=9).value
