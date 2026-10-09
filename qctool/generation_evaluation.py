"""Offline evaluation of generation/runtime artifacts; never model/browser work.

Labels are optional reviewed step expectations. Unlabelled semantic correctness
and false-fail rates are unknown, never reported as zero.
"""
from collections import Counter
from dataclasses import dataclass
import json


def _obligations(report):
    return {o['id']:o for r in report.get('requirements',[]) for o in r.get('obligations',[])}


def _steps(case):
    return [{'action':s['action'],'args':s['args']} for s in case.get('steps',[])]


@dataclass(frozen=True)
class GenerationEvaluationReport:
    metrics: dict
    cases: tuple

    def summary(self):
        return dict(self.metrics)

    def to_dict(self):
        return {'summary':self.metrics,'cases':list(self.cases)}


def evaluate_generation(generation, *, execution=None, ground_truth=None):
    """Consume native artifacts, preserving source and authoring audit history."""
    gen=generation.to_dict() if hasattr(generation,'to_dict') else generation
    runtime=execution.to_dict() if hasattr(execution,'to_dict') else execution or {}
    labels=ground_truth or {}
    obligations=_obligations(gen.get('coverage',{}))
    cases=gen.get('cases',[])
    executable_ids={c['id'] for c in gen.get('executable_cases',[])}
    statuses={c['case_id']:c['status'] for c in runtime.get('results',runtime.get('cases',[]))}
    history={entry['case']['id']:entry for entry in gen.get('agent_history',[])}
    records=list(gen.get('design_audit',[]))+[a for h in history.values() for a in h.get('provider_attempts',[])]
    semantics=Counter();unsafe=0;false_fail=0;false_fail_known=0;review=Counter();rows=[]
    for case in cases:
        label=labels.get(case['id']) or labels.get((case.get('obligation_ids') or [''])[0])
        outcome='unknown'
        if label and 'steps' in label:
            expected=label['steps']
            outcome='correct' if _steps(case)==expected else 'incorrect'
            if outcome=='incorrect' and case['id'] in executable_ids:unsafe+=1
            if statuses.get(case['id'])=='Fail':
                false_fail_known+=1;false_fail+=int(outcome=='incorrect')
        elif label and label.get('expected_abstention'):
            outcome='correct' if case['id'] not in executable_ids else 'incorrect'
            unsafe+=int(outcome=='incorrect')
        semantics[outcome]+=1
        decision=(label or {}).get('review_decision')
        if decision in {'accepted','rejected'}:review[decision]+=1
        rows.append({'testcase_id':case['id'],'obligation_ids':case.get('obligation_ids',[]),
            'generation_status':case.get('generation_status'),'runnable':case['id'] in executable_ids,
            'runtime_result':statuses.get(case['id'],'Not Executed'),'semantic':outcome,
            'source_intent':case.get('source_intent',{}),'history':history.get(case['id'],{})})
    usage=[r.get('usage',{}) for r in records]
    scoped = len({o for c in gen.get('generated_cases',[]) for o in c.get('obligation_ids',[])})
    metrics={'requirements':len(gen.get('coverage',{}).get('requirements',[])),
        'obligations':len(obligations),'selected_obligations':scoped,'runnable_cases':len(executable_ids),'generated_cases':len(gen.get('generated_cases',[])),
        'executable_coverage':sum(o['status']=='covered' for o in obligations.values()),
        'runtime_results':dict(Counter(statuses.get(c['id'],'Not Executed') for c in cases)),
        'semantic':dict(semantics),'unsafe_applied_repairs':unsafe if labels else None,
        'false_fail_rate':false_fail/false_fail_known if false_fail_known else None,
        'false_fail_labelled_cases':false_fail_known,
        'review_acceptance_rate':review['accepted']/sum(review.values()) if review else None,
        'calls':len(records),'unknown_usage_calls':sum(not u for u in usage),
        'prompt_tokens':sum(u.get('prompt_tokens',0) for u in usage),
        'completion_tokens':sum(u.get('completion_tokens',0) for u in usage),
        'total_tokens':sum(u.get('total_tokens',0) for u in usage),
        'cache_hit_tokens':sum(u.get('prompt_cache_hit_tokens',0) for u in usage),
        'tokens_per_obligation':sum(u.get('total_tokens',0) for u in usage)/scoped if scoped else None,
        'calls_per_obligation':len(records)/scoped if scoped else None,
        'generation_status':gen.get('status')}
    return GenerationEvaluationReport(metrics,tuple(rows))


def read_review_labels(path):
    """Read QC decisions without changing runtime status or expected outcomes."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if 'Review' not in wb.sheetnames:
            return {}
        return {r[0]: {'review_decision': r[3]} for r in wb['Review'].iter_rows(min_row=2, values_only=True)
                if r[0] and r[3] in {'accepted', 'rejected'}}
    finally:
        wb.close()
