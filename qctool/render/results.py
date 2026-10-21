"""Ghi kết quả chạy ngược vào file testcase theo template.

Đây là lý do 14 cột template có 4 cột runtime (`Actual Result`, `Status`,
`Environment`, `Executed By / Date`): file team nhận được là **bản ghi hai pha** —
sinh ra một lần, rồi chạy điền vào chính file đó.

Nhờ vậy QC không phải đối chiếu hai file với nhau.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from qctool.execute.runner import RunResult

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
GREY = PatternFill("solid", fgColor="E7E6E6")

STATUS_FILL = {
    "Pass": GREEN,
    "Fail": RED,
    "Blocked": AMBER,
    "Not Executed": GREY,
}

# Vị trí cột theo template 14 cột (1-based)
COL_ACTUAL = 9
COL_STATUS = 10
COL_ENV = 12
COL_EXECUTED = 14


def _evidence_note(evidence: dict) -> str:
    """Dòng trỏ tới bằng chứng, ghi ngay dưới diễn biến trong cột Actual Result."""
    if not evidence:
        return ""
    parts: list[str] = []
    for shot in evidence.get("screenshots", []) or []:
        parts.append(f"ảnh: evidence/{shot}")
    if evidence.get("trace"):
        parts.append(f"trace: evidence/{evidence['trace']}")
    if count := len(evidence.get("console") or []):
        parts.append(f"{count} lỗi JS")
    if count := len(evidence.get("network") or []):
        parts.append(f"{count} request lỗi")
    if count := len(evidence.get("axe") or []):
        parts.append(f"{count} vi phạm accessibility")
    return ("Bằng chứng: " + " · ".join(parts)) if parts else ""


def fill_results(xlsx_path: Path, result: RunResult, *, executed_by: str = "qc-tool") -> int:
    """Điền kết quả vào file đã sinh. Trả số dòng đã cập nhật."""
    if not xlsx_path.is_file():
        raise FileNotFoundError(
            f"Không tìm thấy {xlsx_path}. Phải sinh file testcase trước khi chạy."
        )

    wb = load_workbook(xlsx_path)
    ws = wb["Test Cases"] if "Test Cases" in wb.sheetnames else wb.active

    by_id = {c.case_id: c for c in result.cases}
    updated = 0

    for row in range(2, ws.max_row + 1):
        case_id = ws.cell(row=row, column=1).value
        if not case_id:
            continue
        case = by_id.get(str(case_id).strip())
        if case is None:
            continue

        actual = "\n".join(case.actual) if case.actual else ""
        if case.flaky:
            actual = ("⚠ Test chập chờn: lần chạy đầu không đạt, chạy lại mới đạt.\n"
                      + actual)
        if case.blocked_reason:
            actual = (actual + "\n" if actual else "") + f"Lý do chặn: {case.blocked_reason}"
        if paths := _evidence_note(case.evidence):
            actual = (actual + "\n" if actual else "") + paths

        ws.cell(row=row, column=COL_ACTUAL, value=actual)
        status_cell = ws.cell(row=row, column=COL_STATUS, value=case.status)
        status_cell.fill = STATUS_FILL.get(case.status, GREY)
        status_cell.font = Font(size=10, bold=True)
        ws.cell(row=row, column=COL_ENV, value=result.environment or result.base_url)
        ws.cell(row=row, column=COL_EXECUTED,
                value=f"{executed_by} · {result.finished_at}")

        for col in (COL_ACTUAL, COL_ENV, COL_EXECUTED):
            ws.cell(row=row, column=col).alignment = Alignment(vertical="top", wrap_text=True)

        updated += 1

    wb.save(xlsx_path)
    return updated
