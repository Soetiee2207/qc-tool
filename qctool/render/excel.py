"""Render testcase IR → file Excel theo template của team.

IR là trung tâm; file này chỉ là một output adapter. Đổi template thì sửa ở đây,
không đụng tới generator.

Cấu trúc cột lấy **nguyên văn** từ ảnh spec "Standard Test Case Template Structure"
(14 cột, xem `templates/testcase/README.md`).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..testcase import TestCase

# (tên cột, bề rộng, có phải cột runtime không)
COLUMNS = [
    ("Test Case ID", 15, False),
    ("Module / Feature Name", 22, False),
    ("Test Title / Summary", 46, False),
    ("Description / Objective", 34, False),
    ("Preconditions", 30, False),
    ("Test Data", 26, False),
    ("Test Steps", 52, False),
    ("Expected Result", 40, False),
    ("Actual Result", 30, True),
    ("Status", 14, True),
    ("Priority / Severity", 14, False),
    ("Environment", 20, True),
    ("Created By / Date", 18, False),
    ("Executed By / Date", 18, True),
]

STATUS_VALUES = ["Pass", "Fail", "Blocked", "Not Executed"]
PRIORITY_VALUES = ["High", "Medium", "Low"]

GEN_HDR = PatternFill("solid", fgColor="1F4E79")
RUN_HDR = PatternFill("solid", fgColor="9C5700")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
BODY = Font(size=10)
ID_FONT = Font(size=10, bold=True, color="1F4E79")

THIN = Side(style="thin", color="D1D1D1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
GREY = PatternFill("solid", fgColor="E7E6E6")


def render(cases: list[TestCase], out_path: Path, *, created_by: str = "qc-tool",
           environment: str = "") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    _header(ws)
    row = 2
    for case in cases:
        _write_case(ws, row, case, created_by, environment)
        row += 1

    _decorate(ws, last_row=row - 1)
    agent_cases = [c for c in cases if c.generation_status in {'ready', 'needs_review', 'manual', 'blocked_env'}]
    if agent_cases:
        review = wb.create_sheet('Review')
        review.append(['Test Case ID', 'Generation status', 'Review reasons', 'Decision', 'Reviewer note'])
        decision = DataValidation(type='list', formula1='"accepted,rejected"', allow_blank=True)
        review.add_data_validation(decision)
        for case in agent_cases:
            review.append([case.id, case.generation_status, '; '.join(case.review_reasons), None, None])
        decision.add(f'D2:D{len(agent_cases)+1}')
        review.freeze_panes = 'A2'
        for column, width in {'A':24,'B':20,'C':75,'D':18,'E':65}.items():
            review.column_dimensions[column].width = width
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def _header(ws) -> None:
    for idx, (name, width, is_run) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
        cell = ws.cell(row=1, column=idx, value=name)
        cell.fill = RUN_HDR if is_run else GEN_HDR
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 34


def _write_case(ws, row: int, case: TestCase, created_by: str, environment: str) -> None:
    values = [
        case.id,
        case.module,
        case.title,
        _objective_text(case),
        "\n".join(f"- {p}" for p in case.preconditions),
        _test_data_text(case),
        case.steps_text(),
        case.expected_text(),
        "",                       # Actual Result — điền khi chạy
        None,                     # Status
        case.priority.value,
        environment,              # Environment — điền khi chạy
        f"{created_by}" + (f" · {case.created_at}" if case.created_at else ""),
        "",                       # Executed By / Date
    ]
    for idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=idx, value=value)
        cell.font = ID_FONT if idx == 1 else BODY
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BORDER

    ws.row_dimensions[row].height = _row_height(case)


def _objective_text(case: TestCase) -> str:
    """Cột Description / Objective — có đánh dấu case thủ công.

    Template 14 cột không có cột riêng cho "máy chạy hay người chạy", mà đây là
    thông tin người đọc cần biết trước tiên. Đánh dấu ngay đầu cột mô tả thay vì
    thêm cột mới — cấu trúc template là hợp đồng với người dùng, không tự đổi.
    """
    if case.generation_status == 'needs_review':
        return f"[Needs review] {case.objective}\n" + '; '.join(case.review_reasons)
    if case.execution == "manual":
        return f"[Thủ công] {case.objective}"
    return case.objective


def _test_data_text(case: TestCase) -> str:
    if not case.test_data:
        return ""
    lines = []
    for key, value in case.test_data.items():
        if isinstance(value, (dict, list)):
            import json

            value = json.dumps(value, ensure_ascii=False)
        text = str(value)
        if len(text) > 120:
            text = text[:117] + "…"
        lines.append(f"{key}: {text}")
    return "\n".join(lines)


def _row_height(case: TestCase) -> int:
    """Ước lượng chiều cao hàng để nội dung không bị cắt."""
    steps = len(case.steps_text().splitlines())
    expected = len(case.expected_text().splitlines())
    lines = max(steps, expected, 2)
    return min(15 * lines + 8, 320)


def _decorate(ws, last_row: int) -> None:
    end = max(last_row + 1, 2)

    status_dv = DataValidation(type="list", formula1=f'"{",".join(STATUS_VALUES)}"',
                               allow_blank=True, showDropDown=False)
    status_dv.errorTitle = "Giá trị không hợp lệ"
    status_dv.error = "Chỉ được chọn: " + " / ".join(STATUS_VALUES)
    ws.add_data_validation(status_dv)
    status_dv.add(f"J2:J{max(end, 2000)}")

    prio_dv = DataValidation(type="list", formula1=f'"{",".join(PRIORITY_VALUES)}"',
                             allow_blank=True, showDropDown=False)
    ws.add_data_validation(prio_dv)
    prio_dv.add(f"K2:K{max(end, 2000)}")

    rng = f"J2:J{max(end, 2000)}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pass"'], fill=GREEN))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fail"'], fill=RED))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Blocked"'], fill=AMBER))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Not Executed"'], fill=GREY))

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(end, 2000)}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "1F4E79"


def read_review_labels(path):
    """Read QC decisions without changing runtime status or expected outcomes."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if 'Review' not in wb.sheetnames:
            return {}
        return {r[0]: {'review_decision': r[3]} for r in wb['Review'].iter_rows(min_row=2, values_only=True)
                if r[0] and r[3] in {'accepted', 'rejected'}}
    finally:
        wb.close()
